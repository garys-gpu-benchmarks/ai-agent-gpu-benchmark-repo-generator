#!/usr/bin/env python3
"""Ubuntu 24.04 and 26.04 siblings run the same test (TEMPLATE_00_116).

TEMPLATE_00_115's spreadsheet gave NVIDIA 2xx and 4xx different baseline and
extended values in 21 rows (for example x29 output_len 18600 vs 12300), and
24.04 installed torch 2.7.1+cu128 while 26.04 installed 2.13.0+cu130, so the
two NVIDIA result sets were not comparable. AMD 1xx/3xx already matched. The
software inventory also ran the system python3, so it never recorded the
workload's PyTorch version.
"""

from __future__ import annotations

import re
import unittest
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MATRIX = ROOT / "BenchmarkSpecDefinitions.xlsx"
# The only parameters allowed to differ between a 24.04 workload and its
# 26.04 sibling: the expected OS kernel and CUDA toolkit in x01.
ALLOWED = {("01", "kernel_version"), ("01", "cuda_version")}


def _norm(value) -> str:
    text = "" if value is None else str(value).strip()
    try:
        number = float(text)
    except ValueError:
        return text
    return str(int(number)) if number.is_integer() else text


class ParameterParityTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        workbook = load_workbook(MATRIX, read_only=True, data_only=True)
        cls.params = defaultdict(dict)
        for row in list(workbook["Parameters_SmokeBaselineExtend"].iter_rows(values_only=True))[1:]:
            if row[0] is None:
                continue
            cls.params[_norm(row[0])][str(row[2])] = tuple(_norm(v) for v in row[3:6])
        cls.defs = {
            _norm(row[0]): row
            for row in list(workbook["Workload_Definitions"].iter_rows(values_only=True))[1:]
            if row[0] is not None
        }
        header = next(workbook["Workload_Definitions"].iter_rows(values_only=True))
        cls.col = {name: index for index, name in enumerate(header) if name}

    def differences(self, first: str, second: str) -> list[tuple[str, str]]:
        found = []
        for number in range(1, 33):
            suffix = f"{number:02d}"
            a = self.params[f"{first}{suffix}"]
            b = self.params[f"{second}{suffix}"]
            for name in sorted(set(a) | set(b)):
                if a.get(name) != b.get(name) and (suffix, name) not in ALLOWED:
                    found.append((suffix, name))
        return found

    def test_nvidia_2404_and_2604_share_every_profile_value(self) -> None:
        self.assertEqual(self.differences("2", "4"), [])

    def test_amd_2404_and_2604_share_every_profile_value(self) -> None:
        self.assertEqual(self.differences("1", "3"), [])

    def test_nvidia_siblings_list_the_same_pytorch_and_torchvision(self) -> None:
        torch_col = self.col["PyTorch Version"]
        vision_col = self.col["torchvision Version"]
        for number in range(1, 33):
            a = self.defs[f"2{number:02d}"]
            b = self.defs[f"4{number:02d}"]
            with self.subTest(workload=f"x{number:02d}"):
                self.assertEqual(a[torch_col], b[torch_col])
                self.assertEqual(a[vision_col], b[vision_col])


class LinpackSizeTest(unittest.TestCase):
    """The ledger records the spreadsheet's problem_size_N, so it must be the N that runs.

    Until TEMPLATE_00_116, 220/420 extended said 65536 while the collector clamped
    N >= 65536 to 32768, so every extended row recorded a size that never ran.
    """

    def test_spreadsheet_n_is_never_clamped_by_the_collector(self) -> None:
        import importlib.util

        path = ROOT / "implementation_components" / "linpack-hpl-nvidia" / "files" / "scripts" / "collect_workload.py"
        spec = importlib.util.spec_from_file_location("collect_hpl_nvidia", path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        workbook = load_workbook(MATRIX, read_only=True, data_only=True)
        rows = list(workbook["Parameters_SmokeBaselineExtend"].iter_rows(values_only=True))[1:]
        checked = 0
        for row in rows:
            if _norm(row[0]) in ("220", "420") and row[2] == "problem_size_N":
                for profile, value in zip(("smoke", "baseline", "extended"), row[3:6]):
                    n = int(float(value))
                    with self.subTest(workload=_norm(row[0]), profile=profile):
                        self.assertEqual(module.clamp_problem_n(n), n)
                    checked += 1
        self.assertEqual(checked, 6)


class OnePyTorchBuildTest(unittest.TestCase):
    def test_no_nvidia_installer_pins_the_old_24_04_torch(self) -> None:
        paths = [ROOT / "scripts" / "install_pytorch_nvidia.sh", ROOT / "scripts" / "templates" / "setup_nvidia_skeleton.sh"]
        paths += sorted((ROOT / "implementation_components").glob("*-nvidia/files/**/*.sh"))
        for path in paths:
            code = "\n".join(
                line for line in path.read_text(encoding="utf-8").splitlines() if not line.lstrip().startswith("#")
            )
            with self.subTest(path=path.relative_to(ROOT).as_posix()):
                self.assertIsNone(re.search(r"2\.7\.1\+cu128|0\.22\.1\+cu128|whl/cu128", code))

    def test_24_04_default_torch_is_the_26_04_build(self) -> None:
        text = (ROOT / "scripts" / "install_pytorch_nvidia.sh").read_text(encoding="utf-8")
        self.assertEqual(text.count('TORCH_SPEC="${PYTORCH_INSTALL_SPEC:-torch==2.13.0+cu130}"'), 2)
        self.assertEqual(text.count('VISION_SPEC="${TORCHVISION_INSTALL_SPEC:-torchvision==0.28.0+cu130}"'), 2)


class InventoryUsesVenvTest(unittest.TestCase):
    def test_generated_collector_puts_the_workload_venv_first(self) -> None:
        for name in ("generate_collect_hw_sw_info.sh", "collect_hw_sw_info.sh"):
            text = (ROOT / "scripts" / name).read_text(encoding="utf-8")
            with self.subTest(script=name):
                self.assertIn('-x "$COLLECTOR_REPO_ROOT/.venv/bin/python3"', text)
                self.assertIn('export PATH="$VIRTUAL_ENV/bin:$PATH"', text)
        collector = (ROOT / "scripts" / "collect_hw_sw_info.sh").read_text(encoding="utf-8")
        self.assertLess(collector.index('export PATH="$VIRTUAL_ENV/bin:$PATH"'), collector.index("detect_vendor()"))


if __name__ == "__main__":
    unittest.main()
