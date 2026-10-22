#!/usr/bin/env python3
"""TEMPLATE_00_117 fixes from the TEMPLATE_00_116 Cursor session reviews.

1. Workbook authoring notes no longer fail generation. In the 201-232 and
   401-432 batches, 207 and 416 failed self_check because a
   "TEMPLATE_00_116" note in Parameters_SmokeBaselineExtend was copied into
   benchmark_specification.json. Every profile cell is now cleaned.
2. x09/x10 smoke DRAM tiers are larger than the host LLC (they were 16 MiB,
   inside a 48 MiB L3).
3. x14 smoke max_bytes / num_iterations in the workbook are what the
   collectors run (the collectors raised 4096 bytes / 3 iterations to their
   256 MiB / 10 floors without saying so).
4. x01 NVIDIA driver compliance: the workbook describes the minimum-version
   rule the collector applies, not a substring match.
"""

from __future__ import annotations

import importlib.util
import re
import sys
import unittest
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
COMPONENTS = ROOT / "implementation_components"
MATRIX = ROOT / "BenchmarkSpecDefinitions.xlsx"
ALL_WORKLOADS = [str(n) for base in (100, 200, 300, 400) for n in range(base + 1, base + 33)]
MIB = 1024**2
DRAM_MINIMUM = 256 * MIB


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


EXTRACT = _load(ROOT / "scripts" / "extract_benchmark_definition.py", "extract_00_117_review")
SHEETS = EXTRACT._load_workbook(MATRIX)


def profile(workload: str) -> dict[str, dict[str, str]]:
    return EXTRACT._extract_profile_parameter_values(SHEETS, workload)[1]


class TemplateNoteStripTest(unittest.TestCase):
    def test_dated_notes_lose_the_tag(self) -> None:
        for text in (
            "2026-10-22 TEMPLATE_00_117: smoke array_size 1048576 -> 67108864. Smoke fits in L2.",
            "Sequential prompt-response: 8 prompts. | 2026-10-22 TEMPLATE_00_117: baseline 8 -> 5.",
            "TEMPLATE_00_116 2026-10-21: extended 65536 -> 32768",
            "kept (see TEMPLATE_00_90 retune) for parity",
        ):
            with self.subTest(text=text):
                self.assertNotIn("TEMPLATE_00_", EXTRACT._strip_profile_cell(text).upper())

    def test_untagged_cells_are_returned_unchanged(self) -> None:
        for text in ("4096", "0.5", ".5", "l1=16384,l2=262144", "a  b.", "", "1e-3;"):
            with self.subTest(text=text):
                self.assertEqual(EXTRACT._strip_profile_cell(text), text)

    def test_no_workload_extracts_an_authoring_tag(self) -> None:
        for workload in ALL_WORKLOADS:
            values = profile(workload)
            with self.subTest(workload=workload):
                self.assertTrue(values)
                for name, row in values.items():
                    for column, cell in row.items():
                        self.assertNotIn("TEMPLATE_00_", cell.upper(), f"{workload} {name} {column}")

    def test_workbook_still_has_tagged_notes_so_the_strip_is_exercised(self) -> None:
        workbook = load_workbook(MATRIX, read_only=True)
        notes = [row[6] for row in workbook["Parameters_SmokeBaselineExtend"].iter_rows(min_row=2, values_only=True)]
        self.assertTrue(any(isinstance(n, str) and "TEMPLATE_00_" in n for n in notes))


class SmokeDramTierTest(unittest.TestCase):
    def test_x09_smoke_dram_tiers_exceed_llc(self) -> None:
        for workload in ("109", "209", "309", "409"):
            tiers = dict(item.split("=") for item in profile(workload)["working_set_size"]["smoke"].split(","))
            with self.subTest(workload=workload):
                self.assertGreaterEqual(int(tiers["local_dram"]), DRAM_MINIMUM)
                self.assertGreaterEqual(int(tiers["remote_dram"]), DRAM_MINIMUM)

    def test_x10_smoke_dram_tier_exceeds_llc(self) -> None:
        for workload in ("110", "210", "310", "410"):
            row = profile(workload)
            sizes = [int(s) for s in row["working_set_size"]["smoke"].split(",")]
            levels = row["cache_level"]["smoke"].split(",")
            with self.subTest(workload=workload):
                self.assertEqual(len(sizes), len(levels))
                self.assertGreaterEqual(sizes[levels.index("DRAM")], DRAM_MINIMUM)


def _floor(path: Path, name: str) -> int:
    match = re.search(rf"(?m)^{name} = (.+)$", path.read_text(encoding="utf-8"))
    assert match, f"{name} missing from {path}"
    return int(eval(match.group(1), {"__builtins__": {}}))


X14_COLLECTORS = {
    "amd": COMPONENTS / "hipmemcpy-bandwidth-amd" / "files" / "scripts" / "collect_workload.py",
    "nvidia": COMPONENTS / "cuda-memcpy-nvidia" / "files" / "scripts" / "collect_workload.py",
}


class X14WorkbookMatchesCollectorTest(unittest.TestCase):
    def test_every_profile_meets_the_collector_floors(self) -> None:
        for workload in ("114", "214", "314", "414"):
            collector = X14_COLLECTORS["nvidia" if workload[0] in "24" else "amd"]
            size_floor = _floor(collector, "BANDWIDTH_SIZE_FLOOR")
            iteration_floor = _floor(collector, "BANDWIDTH_ITERATION_FLOOR")
            row = profile(workload)
            for column in ("smoke", "baseline", "extended"):
                with self.subTest(workload=workload, profile=column):
                    self.assertGreaterEqual(int(row["max_bytes"][column]), size_floor)
                    self.assertGreaterEqual(int(row["num_iterations"][column]), iteration_floor)

    def test_smoke_values(self) -> None:
        for workload in ("114", "214", "314", "414"):
            row = profile(workload)
            with self.subTest(workload=workload):
                self.assertEqual(row["max_bytes"]["smoke"], "268435456")
                self.assertEqual(row["num_iterations"]["smoke"], "10")


class X01DriverComplianceWordingTest(unittest.TestCase):
    def test_workbook_describes_the_minimum_rule(self) -> None:
        workbook = load_workbook(MATRIX, read_only=True)
        rows = list(workbook["Workload_Definitions"].iter_rows(values_only=True))
        col = rows[0].index("Execution Description With Parameters")
        seen = 0
        for row in rows[1:]:
            if str(row[0]) in ("201", "401"):
                seen += 1
                text = row[col]
                with self.subTest(workload=row[0]):
                    self.assertNotIn("substring of nvidia-smi Driver Version", text)
                    self.assertIn("nvidia_driver_version: minimum driver version", text)
                    self.assertIn("at or above", text)
        self.assertEqual(seen, 2)

    def test_parameter_reference_matches(self) -> None:
        params = yaml.safe_load((ROOT / "config" / "workload_parameters.yaml").read_text(encoding="utf-8"))
        text = yaml.safe_dump(params)
        self.assertNotIn("substring of nvidia-smi Driver Version", text)

    def test_collector_scores_at_or_above_the_minimum(self) -> None:
        collector = _load(
            COMPONENTS / "system-config-nvidia" / "files" / "scripts" / "collect_workload.py", "sysconfig_nv_00_117"
        )
        self.assertTrue(collector.version_at_least("595.91.07", "580"))
        self.assertTrue(collector.version_at_least("580.0", "580"))
        self.assertFalse(collector.version_at_least("570.172.08", "580"))


if __name__ == "__main__":
    unittest.main()
