#!/usr/bin/env python3
"""Regression checks for the 2026-10-18 metric key and label renames.

Each renamed numbered metric must resolve, from the workbook's Metrics text,
to the key its collector now emits; summary.json files written before the
rename must still resolve through STORAGE_ALIASES; and the collectors and
gates must use the new keys.
"""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from metric_contract import (  # noqa: E402
    STORAGE_ALIASES,
    apply_storage_aliases,
    classify_na,
    publication_status,
)
from print_metric_summary import extract_paren_keys, metric_items, resolve_metrics  # noqa: E402

IC = ROOT / "implementation_components"

# workload -> {metric number: key the collector emits}
RENAMED = {
    101: {"2": "kernel_release_mismatch_count"},
    301: {"2": "kernel_release_mismatch_count"},
    103: {"3": "ras_uncorrectable_ecc_delta_count", "4": "thermal_throttle_event_count"},
    303: {"3": "ras_uncorrectable_ecc_delta_count", "4": "thermal_throttle_event_count"},
    104: {"5": "loop_verification_pass_pct"},
    304: {"5": "loop_verification_pass_pct"},
    109: {"5": "latency_ns_remote_dram"},
    116: {"1": "allreduce_busbw_gb_s", "2": "allreduce_algbw_gb_s", "3": "allreduce_latency_at_min_bytes_us"},
    316: {"1": "allreduce_busbw_gb_s", "2": "allreduce_algbw_gb_s", "3": "allreduce_latency_at_min_bytes_us"},
    120: {"1": "sustained_fp64_performance_across_problem_size_n_tflops"},
    203: {"1": "peak_gpu_temperature_c"},
    403: {"1": "peak_gpu_temperature_c"},
    204: {"4": "hbm3_coverage_gb", "5": "fixed_pattern_fill_verify_pass_rate"},
    404: {"4": "hbm3_coverage_gb", "5": "fixed_pattern_fill_verify_pass_rate"},
    216: {"1": "allreduce_busbw_gb_s", "2": "allreduce_algbw_gb_s", "3": "allreduce_latency_at_min_bytes_us"},
    416: {"1": "allreduce_busbw_gb_s", "2": "allreduce_algbw_gb_s", "3": "allreduce_latency_at_min_bytes_us"},
    218: {"2": "tflops_fwd", "3": "min_operand_traffic_gb_s"},
    418: {"2": "tflops_fwd", "3": "min_operand_traffic_gb_s"},
    221: {"3": "fp16_tflops"},
    227: {"1": "decode_throughput_tokens_sec"},
    427: {"1": "decode_throughput_tokens_sec"},
}

# old stored key -> new numbered key (summary.json written before the rename)
OLD_TO_NEW = {
    "kernel_driver_mismatch_count": "kernel_release_mismatch_count",
    "even_word_store_pass_rate": "loop_verification_pass_pct",
    "algbw_gb_s": "allreduce_algbw_gb_s",
    "busbw_gb_s": "allreduce_busbw_gb_s",
    "walking_1s_moving_inversion_pass_rate": "fixed_pattern_fill_verify_pass_rate",
    "sustained_decode_throughput_under_cache_pressure_tokens_sec": "decode_throughput_tokens_sec",
}


def workbook_metrics() -> dict[int, str]:
    from openpyxl import load_workbook

    sheet = load_workbook(ROOT / "BenchmarkSpecDefinitions.xlsx", read_only=True)["Workload_Definitions"]
    rows = list(sheet.iter_rows(values_only=True))
    header = list(rows[0])
    number_col, metrics_col = header.index("Workload Number"), header.index("Metrics")
    out = {}
    for row in rows[1:]:
        try:
            out[int(row[number_col])] = str(row[metrics_col] or "")
        except (TypeError, ValueError):
            continue
    return out


def items_for(metrics_text: str) -> list[tuple[str, str]]:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "benchmark_specification.json"
        path.write_text(json.dumps([{"field_name": "Metrics", "value": metrics_text}]), encoding="utf-8")
        return metric_items(path)


class WorkbookResolutionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.metrics = workbook_metrics()

    def test_renamed_metrics_name_the_emitted_key_first(self):
        for workload, expected in RENAMED.items():
            items = dict(items_for(self.metrics[workload]))
            for number, key in expected.items():
                with self.subTest(workload=workload, metric=number):
                    self.assertEqual(extract_paren_keys(items[number])[0], key)

    def test_new_keys_resolve_without_positional_fallback(self):
        for workload, expected in RENAMED.items():
            items = items_for(self.metrics[workload])
            summary = {extract_paren_keys(d)[0]: float(i + 1) for i, (_n, d) in enumerate(items) if extract_paren_keys(d)}
            resolved = resolve_metrics(items, summary)
            by_number = {row[0]: row for row in resolved}  # (number, description, key, value)
            for number, key in expected.items():
                with self.subTest(workload=workload, metric=number):
                    self.assertEqual(by_number[number][2], key)

    def test_old_summary_keys_still_resolve(self):
        cases = {101: "2", 104: "5", 116: "1", 204: "5", 227: "1"}
        for workload, number in cases.items():
            items = items_for(self.metrics[workload])
            new_key = RENAMED[workload][number]
            old_key = next(old for old, new in OLD_TO_NEW.items() if new == new_key)
            with self.subTest(workload=workload):
                resolved = {row[0]: (row[2], row[3]) for row in resolve_metrics(items, {old_key: 7.0})}
                self.assertEqual(resolved[number], (new_key, 7.0))

    def test_rocsolver_label_is_kept(self):
        for workload in (120, 320):
            self.assertIn("rocSOLVER LU", self.metrics[workload])
            self.assertNotIn("rocHPL FP64, TFLOPS", self.metrics[workload])


class ContractTests(unittest.TestCase):
    def test_aliases_map_old_keys_to_new(self):
        for old, new in OLD_TO_NEW.items():
            self.assertEqual(STORAGE_ALIASES[old][0], new)
            self.assertEqual(apply_storage_aliases({old: 3.0})[new], 3.0)

    def test_sibling_canonical_keys_are_not_aliased(self):
        for key in ("gpu_system_ecc_error_count", "peak_gpu_junction_temp_c", "latency_us", "fp16_tflops", "tflops_fwd"):
            self.assertNotIn(key, STORAGE_ALIASES)

    def test_throttle_not_exposed_is_expected_na(self):
        value = "na (expected: throttle counter not exposed)"
        self.assertEqual(classify_na("thermal_throttle_event_count", value, {}), "expected")
        status, _count = publication_status([("thermal_throttle_event_count", value), ("sustained_gpu_power_w", 300.0)])
        self.assertEqual(status, "N/A-EXPECTED")

    def test_missing_llc_load_counters_are_expected_na(self):
        value = "na (LLC load counters unavailable)"
        self.assertEqual(classify_na("l3_miss_rate", value, {}), "expected")
        status, _count = publication_status([("l3_miss_rate", value), ("giga_updates_score", 0.01)])
        self.assertEqual(status, "N/A-EXPECTED")

    def test_single_gpu_allreduce_latency_is_expected_na(self):
        self.assertEqual(
            classify_na("allreduce_latency_at_min_bytes_us", "na (not measured)", {"gpu_count": 1}),
            "expected",
        )


class CollectorTests(unittest.TestCase):
    def read(self, rel: str) -> str:
        return (IC / rel).read_text(encoding="utf-8")

    def test_collectors_emit_new_keys(self):
        cases = {
            "system-config-amd/files/scripts/collect_workload.py": ["kernel_release_mismatch_count"],
            "system-stress-amd/files/scripts/collect_workload.py": ["ras_uncorrectable_ecc_delta_count", "na (expected: throttle counter not exposed)"],
            "sdc-ecc-amd/files/scripts/collect_workload.py": ["loop_verification_pass_pct"],
            "rccl-bandwidth-amd/files/scripts/collect_workload.py": ["allreduce_busbw_gb_s", "allreduce_algbw_gb_s", "allreduce_latency_at_min_bytes_us"],
            "system-stress-nvidia/files/scripts/collect_workload.py": ["peak_gpu_temperature_c"],
            "sdc-ecc-nvidia/files/scripts/collect_workload.py": ["fixed_pattern_fill_verify_pass_rate"],
            "nccl-bandwidth-nvidia/files/scripts/collect_workload.py": ["allreduce_busbw_gb_s", "allreduce_algbw_gb_s", "allreduce_latency_at_min_bytes_us"],
            "vllm-kvcache-nvidia/files/scripts/benchmark_serving.py": ["decode_throughput_tokens_sec"],
        }
        stale = ["kernel_driver_mismatch_count", "gpu_system_ecc_error_count\"", "even_word_store_pass_rate",
                 "peak_gpu_junction_temp_c", "walking_1s_moving_inversion_pass_rate",
                 "sustained_decode_throughput_under_cache_pressure_tokens_sec"]
        for rel, keys in cases.items():
            text = self.read(rel)
            for key in keys:
                with self.subTest(file=rel, key=key):
                    self.assertIn(key, text)
            if "system-stress-nvidia" in rel:
                continue  # 203/403 keeps gpu_system_ecc_error_count
            for key in stale:
                if key == "peak_gpu_junction_temp_c" and "system-stress-amd" in rel:
                    continue  # 103/303 reads junction/hotspot and keeps that key
                with self.subTest(file=rel, stale=key):
                    self.assertNotIn(key, text)

    def test_gates_use_new_keys(self):
        config = json.loads(self.read("system-config-amd/gate.json"))["thresholds"]
        stress = json.loads(self.read("system-stress-amd/gate.json"))["thresholds"]
        self.assertIn("kernel_release_mismatch_count_max", config)
        self.assertNotIn("kernel_driver_mismatch_count_max", config)
        self.assertIn("ras_uncorrectable_ecc_delta_count_max", stress)
        self.assertNotIn("gpu_system_ecc_error_count_max", stress)

    def test_nccl_busbw_factors(self):
        text = self.read("nccl-bandwidth-nvidia/files/src/nccl_bw.cu")
        self.assertIn('std::strcmp(name, "all_gather") == 0 || std::strcmp(name, "reduce_scatter") == 0', text)
        self.assertIn("algbw * (double)(ndev - 1) / (double)ndev", text)
        self.assertIn("algbw * 2.0 * (ndev - 1) / (double)ndev", text)

    def test_216_max_bytes_is_one_gib(self):
        from openpyxl import load_workbook

        sheet = load_workbook(ROOT / "BenchmarkSpecDefinitions.xlsx", read_only=True)["Parameters_SmokeBaselineExtend"]
        rows = {(str(r[0]), r[2]): r for r in sheet.iter_rows(values_only=True)}
        for workload in ("216", "416"):
            row = rows[(workload, "max_bytes")]
            self.assertEqual((int(row[4]), int(row[5])), (1073741824, 1073741824))


if __name__ == "__main__":
    unittest.main()
