#!/usr/bin/env python3
"""TEMPLATE_00_117: x30 publishes medians, and every run records its release.

x30 (SGLang prompt-response) averaged 8 timed requests. On NVIDIA 24.04 one
request still waited on compilation after the warm-up (TTFT 9156 ms; the other
seven ~29 ms), so the published TTFT was 1170 ms. The software inventory also
did not record which release a run came from.
"""

from __future__ import annotations

import ast
import base64
import re
import statistics
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
COMPONENTS = ROOT / "implementation_components"
CLIENTS = [COMPONENTS / f"sglang-prompt-response-{v}" / "files" / "scripts" / "prompt_response_client.py" for v in ("amd", "nvidia")]
PARSERS = [
    COMPONENTS / "sglang-prompt-response-nvidia" / "files" / "scripts" / "parse_results.py",
    COMPONENTS / "sglang-prompt-parser-amd" / "files" / "scripts" / "parse_results.py",
]
MATRIX = ROOT / "BenchmarkSpecDefinitions.xlsx"
MEDIAN_KEYS = {"end_to_end_latency_ms", "ttft_ms", "output_tokens_per_s"}


class X30MedianTest(unittest.TestCase):
    def test_clients_summarize_with_the_median(self) -> None:
        for path in CLIENTS:
            text = path.read_text(encoding="utf-8")
            with self.subTest(path=path.relative_to(ROOT).as_posix()):
                self.assertIn("return statistics.median(nums)", text)
                self.assertNotIn("return sum(nums) / len(nums)", text)
                self.assertRegex(text, r"(?m)^import statistics$")

    def test_parsers_use_the_median_for_per_request_metrics_only(self) -> None:
        for path in PARSERS:
            text = path.read_text(encoding="utf-8")
            with self.subTest(path=path.relative_to(ROOT).as_posix()):
                match = re.search(r"MEDIAN_METRICS = (\{[^}]*\})", text)
                self.assertIsNotNone(match)
                self.assertEqual(ast.literal_eval(match.group(1)), MEDIAN_KEYS)
                self.assertIn("summarize = statistics.median if name in MEDIAN_METRICS else statistics.fmean", text)

    def test_one_slow_request_no_longer_dominates(self) -> None:
        ttft = [9156.1, 30.7, 28.7, 30.3, 28.4, 29.0, 29.1, 28.9]  # 230 baseline, v1.0.21
        self.assertGreater(statistics.fmean(ttft), 1000)
        self.assertLess(statistics.median(ttft), 31)

    def test_labels_say_p50(self) -> None:
        workbook = load_workbook(MATRIX, read_only=True, data_only=True)
        rows = list(workbook["Workload_Definitions"].iter_rows(values_only=True))
        col = rows[0].index("Metrics")
        for row in rows[1:]:
            if str(row[0]) in ("130", "230", "330", "430"):
                with self.subTest(workload=row[0]):
                    self.assertIn("#1: E2E latency p50, ms (end_to_end_latency_ms)", row[col])
                    self.assertIn("#2: Time To First Token p50, ms (ttft_ms)", row[col])
                    self.assertIn("#3: Output token generation rate (decode) p50, tokens/s (output_tokens_per_s)", row[col])


class ReleaseInInventoryTest(unittest.TestCase):
    def test_command_list_records_workload_and_bundle_release(self) -> None:
        rows = list(load_workbook(ROOT / "config" / "hw_sw_info_commands.xlsx", read_only=True).worksheets[0].iter_rows(values_only=True))
        by_seq = {str(r[0]): r for r in rows[1:] if r and r[0] is not None}
        self.assertIn("describe --tags", by_seq["139.1"][3])
        self.assertIn('"${COLLECTOR_REPO_ROOT:-.}"', by_seq["139.1"][3])
        self.assertIn('"${COLLECTOR_REPO_ROOT:-.}/.."', by_seq["139.2"][3])
        for seq in ("139.1", "139.2"):
            self.assertEqual(by_seq[seq][1], "AMD, nVidia")
            self.assertEqual(by_seq[seq][2], "software")

    def test_generated_collector_exports_the_repo_root_and_carries_the_commands(self) -> None:
        collector = (ROOT / "scripts" / "collect_hw_sw_info.sh").read_text(encoding="utf-8")
        self.assertIn("export COLLECTOR_REPO_ROOT", collector)
        commands = [base64.b64decode(line.split("\t")[3]).decode() for line in collector.splitlines() if line.startswith(("139.1\t", "139.2\t"))]
        self.assertEqual(len(commands), 2)
        self.assertTrue(all("git -C" in c and "describe --tags --always" in c for c in commands))


if __name__ == "__main__":
    unittest.main()
