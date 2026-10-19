#!/usr/bin/env python3
# File: tests/test_metric_pair_parity.py
# Description: Workloads 1NN/3NN (AMD, Ubuntu 24.04/26.04) and 2NN/4NN (NVIDIA)
#   run the same collector, so their Metrics text must be identical. Fixed
#   2026-10-18 for 117/317, 121/321, 128/328, 132/232/332/432, 209/409,
#   210/410, 212/412, 217/417 and 221/421. The labels changed then must name a
#   key their collector writes.
from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
IC = ROOT / "implementation_components"

# workload -> (metric number, key, collector file that must write the key)
EMITTED = {
    117: ("1", "achieved_compute_tflops", "rocblas-gemm-amd/files/scripts/collect_workload.py"),
    217: ("1", "achieved_compute_tflops", "cublas-gemm-nvidia/files/scripts/collect_cublas_gemm.py"),
    121: ("3", "modeled_tflops", "resnet50-training-amd/files/scripts/collect_workload.py"),
    221: ("3", "fp16_tflops", "resnet50-training-nvidia/files/scripts/train_resnet50.py"),
}


def workbook_rows() -> dict[int, dict[str, str]]:
    sheet = load_workbook(ROOT / "BenchmarkSpecDefinitions.xlsx", read_only=True)["Workload_Definitions"]
    rows = list(sheet.iter_rows(values_only=True))
    header = list(rows[0])
    out = {}
    for row in rows[1:]:
        try:
            out[int(row[0])] = {str(name): str(value or "") for name, value in zip(header, row)}
        except (TypeError, ValueError):
            continue
    return out


class MetricPairParityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.rows = workbook_rows()

    def test_ubuntu_pairs_have_identical_metrics(self) -> None:
        pairs = [(a, a + 200) for a in self.rows if a // 100 in (1, 2) and a + 200 in self.rows]
        self.assertGreater(len(pairs), 60)
        for a, b in pairs:
            with self.subTest(pair=f"{a}/{b}"):
                self.assertEqual(self.rows[a]["Metrics"].strip(), self.rows[b]["Metrics"].strip())

    def test_changed_labels_name_an_emitted_key(self) -> None:
        for workload, (number, key, collector) in EMITTED.items():
            with self.subTest(workload=workload):
                self.assertIn(f"#{number}: ", self.rows[workload]["Metrics"])
                self.assertIn(f"({key})", self.rows[workload]["Metrics"])
                self.assertIn(key, (IC / collector).read_text(encoding="utf-8"))

    def test_121_rows_describe_the_amd_collector(self) -> None:
        for workload in (121, 321):
            row = self.rows[workload]
            for column in ("Output Fields / Metrics to Parse", "Raw Output Example", "Parser and Normalization Notes"):
                with self.subTest(workload=workload, column=column):
                    self.assertNotIn("achieved_TFLOPS", row[column])
                    self.assertNotIn("memory_bandwidth_GBps", row[column])
                    self.assertIn("modeled_tflops", row[column])

    def test_no_fixed_cache_resident_qualifier_in_labels(self) -> None:
        # The cache-resident note is added at run time by metric_contract.annotate_metrics.
        for workload, row in self.rows.items():
            with self.subTest(workload=workload):
                self.assertNotIn("cache-resident below", row["Metrics"])


if __name__ == "__main__":
    unittest.main()
